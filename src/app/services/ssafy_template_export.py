from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from posixpath import normpath
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from app.services.report_summarizer import TeamReport
from app.services.ssafy_reporting import WeeklyReport, TEAM_CODES, WEEKLY_FIELDS


_NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/package/2006/relationships",
}

def _tag(ns: str, name: str) -> str:
    return f"{{{_NS[ns]}}}{name}"


def _validate_path(source: str | Path, dest: str | Path, extension: str) -> tuple[Path, Path]:
    src, dst = Path(source), Path(dest)
    if not src.is_file() or src.suffix.lower() != extension:
        raise ValueError(f"원본 {extension} 템플릿 파일을 선택해 주세요.")
    if src.resolve() == dst.resolve():
        raise ValueError("템플릿 파일은 덮어쓰지 않습니다. 다른 이름으로 저장해 주세요.")
    return src, dst


def export_daily_docx(template: str | Path, output: str | Path, reports: dict[str, TeamReport]) -> list[str]:
    """5반 A501~A509 내용/종합 셀만 변경. 다른 팀/문단/표/서식 유지."""
    from docx import Document

    source, dest = _validate_path(template, output, ".docx")
    doc = Document(source)
    applied = set()
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 3:
                continue
            team = row.cells[0].text.strip()
            key = row.cells[1].text.strip()
            if team not in reports or key not in {"내용", "종합"}:
                continue
            report = reports[team]
            values = report.content if key == "내용" else [report.summary]
            new_content = "\n• ".join(values)
            _replace_cell_text(row.cells[2], new_content)
            applied.add((team, key))
    missing = sorted(team for team in reports if (team, "내용") not in applied or (team, "종합") not in applied)
    if missing:
        raise ValueError("DOCX 템플릿에서 내용/종합 항목을 찾지 못했습니다: " + ", ".join(missing))
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    return sorted(reports)


def _replace_cell_text(cell, value: str) -> None:
    """기존 셀의 첫 문단/서식 속성을 최대한 살리고 텍스트만 교체."""
    para = cell.paragraphs[0]
    runs = para.runs
    if runs:
        runs[0].text = value
        for r in runs[1:]:
            r.text = ""
    else:
        para.add_run(value)
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)


def list_week_sheets(path: str | Path) -> list[str]:
    with ZipFile(path) as zf:
        root = ET.fromstring(zf.read("xl/workbook.xml"))
    return [sh.attrib.get("name", "") for sh in root.findall(".//m:sheet", _NS)
            if re.fullmatch(r"\d+주차 진행현황", sh.attrib.get("name", ""))]


def _sheet_path(zf: ZipFile, sheet_name: str) -> str:
    root = ET.fromstring(zf.read("xl/workbook.xml"))
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    ids = {r.attrib.get("Id"): r.attrib.get("Target") for r in rels.findall("p:Relationship", _NS)}
    for sh in root.findall(".//m:sheet", _NS):
        if sh.attrib.get("name") != sheet_name:
            continue
        rel_id = sh.attrib.get(_tag("r", "id"))
        target = ids.get(rel_id)
        if target is None:
            break
        if target.startswith("/"):
            return target.lstrip("/")
        return normpath("xl/" + target)
    raise ValueError(f"'{sheet_name}' 시트를 찾지 못했습니다.")


def _shared_strings(zf: ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in zf.namelist():
        return []
    root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    return ["".join(t.text or "" for t in si.iter(_tag("m", "t")))
            for si in root.findall("m:si", _NS)]


def _value(cell, strings: list[str]) -> str:
    if cell is None:
        return ""
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        return "".join(t.text or "" for t in cell.iter(_tag("m", "t")))
    v = cell.find("m:v", _NS)
    if v is None:
        return ""
    if cell_type == "s":
        return strings[int(v.text)] if v.text is not None else ""
    return v.text or ""


def _write_inline(cell, text: str) -> None:
    for child in list(cell):
        cell.remove(child)
    cell.set("t", "inlineStr")
    is_el = ET.SubElement(cell, _tag("m", "is"))
    t_el = ET.SubElement(is_el, _tag("m", "t"))
    t_el.text = text
    if text != text.strip():
        t_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def export_weekly_xlsx(
    template: str | Path, output: str | Path, reports: dict[str, WeeklyReport],
    sheet_name: str = "1주차 진행현황",
) -> list[str]:
    """OOXML 셀 패치: 해당 주차 시트의 A501~A509 I:N만 수정. 다른 시트는 그대로 복사."""
    source, dest = _validate_path(template, output, ".xlsx")
    bad = sorted(set(reports) - set(TEAM_CODES))
    if bad:
        raise ValueError("서울 5반 이외 팀의 입력은 허용하지 않습니다: " + ", ".join(bad))
    expected_headers = [label for _, label in WEEKLY_FIELDS]
    ET.register_namespace("", _NS["m"])
    ET.register_namespace("r", _NS["r"])

    with ZipFile(source, "r") as zin:
        target_file = _sheet_path(zin, sheet_name)
        strings = _shared_strings(zin)
        xml_root = ET.fromstring(zin.read(target_file))
        sheet_data = xml_root.find("m:sheetData", _NS)
        if sheet_data is None:
            raise ValueError("주차별 시트에서 데이터 테이블을 찾지 못했습니다.")
        rows = sheet_data.findall("m:row", _NS)
        if not rows:
            raise ValueError("선택한 주차 시트가 비어 있습니다.")
        header = {c.attrib.get("r", "")[:1]: _value(c, strings)
                  for c in rows[0].findall("m:c", _NS)}
        for col, label in zip("IJKLMN", expected_headers):
            if header.get(col, "").strip() != label:
                raise ValueError(f"주간 템플릿의 {col}열 제목이 예상과 다릅니다: {header.get(col)}")

        changed = set()
        for row in rows:
            row_no = row.attrib.get("r", "")
            c_lookup = {c.attrib.get("r", ""): c for c in row.findall("m:c", _NS)}
            team = _value(c_lookup.get(f"B{row_no}"), strings).strip()
            if team not in reports:
                continue
            rep = reports[team]
            for col, (field, _) in zip("IJKLMN", WEEKLY_FIELDS):
                ref = f"{col}{row_no}"
                c = c_lookup.get(ref)
                if c is None:
                    c = ET.SubElement(row, _tag("m", "c"), {"r": ref})
                _write_inline(c, getattr(rep, field))
            changed.add(team)
        missing = sorted(set(reports) - changed)
        if missing:
            raise ValueError("주간 템플릿에서 팀코드를 찾지 못했습니다: " + ", ".join(missing))
        new_xml = ET.tostring(xml_root, encoding="utf-8", xml_declaration=True)
        dest.parent.mkdir(parents=True, exist_ok=True)
        with ZipFile(dest, "w") as zout:
            for file in zin.infolist():
                zout.writestr(file, new_xml if file.filename == target_file else zin.read(file.filename))
    return sorted(changed)
