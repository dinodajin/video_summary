"""SSAFY 서울5반 보고서 데이터 포맷/매핑 회귀 테스트 (외부 LLM 불필요)."""
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from app.services.report_summarizer import TeamReport
from app.services.ssafy_reporting import (
    SsafyReportBuilder, WeeklyReport, parse_daily_text, parse_weekly_text,
)
from app.services.ssafy_template_export import export_daily_docx
from app.config import AppConfig


class TestSsafyReport(unittest.TestCase):
    def test_daily_round_trip(self):
        report = TeamReport("A502", ["명세서 분석함.", "멘토 Q&A 진행함."], "기능 범위 확인 단계.")
        parsed = parse_daily_text(report.render(), "A502")
        self.assertEqual(parsed.content, report.content)
        self.assertEqual(parsed.summary, report.summary)

    def test_weekly_round_trip(self):
        weekly = WeeklyReport("분석 단계임.", "A 필수임.", "설계 완료함.", "B 검토함.", "진척 확인되지 않음.", "보안 제약 있음.")
        self.assertEqual(parse_weekly_text(weekly.render()), weekly)

    def test_summary_uses_both_structures(self):
        config = AppConfig()
        model = SsafyReportBuilder(config)
        model._llm._chat_json = lambda *args: {
            "daily": {"content": ["기획을 진행함.", "테스트를 검토함."], "summary": "기획 단계임."},
            "weekly": {"overall": "기획 초기", "basic_requirements": "필수 A"},
        }
        result = model.summarize("기업 멘토에게 프로젝트 요구사항을 물어봄", "A502")
        self.assertEqual(result.daily.team_code, "A502")
        self.assertEqual(result.weekly.basic_requirements, "필수 A")
        self.assertEqual(result.weekly.issues, "전사에서 확인되지 않음.")

    def test_docx_updates_only_selected_team(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            src, dest = Path(tmp) / "template.docx", Path(tmp) / "out.docx"
            doc = Document()
            table = doc.add_table(rows=1, cols=3)
            for team in ("A501", "A502"):
                for label in ("내용", "종합"):
                    cells = table.add_row().cells
                    cells[0].text = team
                    cells[1].text = label
                    cells[2].text = "해당 사항 없음."
            doc.save(src)
            report = TeamReport("A502", ["멘토링 진행함."], "초기 단계임.")
            self.assertEqual(export_daily_docx(src, dest, {"A502": report}), ["A502"])
            result = Document(dest)
            self.assertEqual(result.tables[0].rows[1].cells[2].text, "해당 사항 없음.")
            self.assertIn("멘토링 진행함.", result.tables[0].rows[3].cells[2].text)

if __name__ == "__main__":
    unittest.main()
